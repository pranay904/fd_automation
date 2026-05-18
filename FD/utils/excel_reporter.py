import pandas as pd
import os
from datetime import datetime
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

_BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
_REPORT_DIR = os.path.join(_BASE_DIR, "reports", "excel")

# Colors
GREEN       = PatternFill("solid", fgColor="C6EFCE")
RED         = PatternFill("solid", fgColor="FFC7CE")
YELLOW      = PatternFill("solid", fgColor="FFEB9C")
BLUE_HEADER = PatternFill("solid", fgColor="1F4E79")
GREY        = PatternFill("solid", fgColor="D9D9D9")
ORANGE      = PatternFill("solid", fgColor="F4B942")

WHITE_FONT  = Font(bold=True, color="FFFFFF", size=10)
BOLD_FONT   = Font(bold=True, size=10)
NORMAL_FONT = Font(size=10)

THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

COLUMNS = [
    "Browser", "Q1", "Q2", "Q3", "Q4", "Q5", "Q6",
    "Overall Status", "No Results", "Product Name", "Price",
    "Step", "Expected", "Actual", "Status", "Error"
]


class ExcelReporter:

    def __init__(self):
        self.rows = []
        os.makedirs(_REPORT_DIR, exist_ok=True)
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        self.path = os.path.join(_REPORT_DIR, f"quiz_results_{ts}.xlsx")

    def add_row(self, browser, answers, validations, overall="", product="", price=""):
        no_results = "YES" if product == "No Results" else "NO"
        base = {
            "Browser":        browser,
            "Q1": answers[0], "Q2": answers[1], "Q3": answers[2],
            "Q4": answers[3], "Q5": answers[4], "Q6": answers[5],
            "Overall Status": overall,
            "No Results":     no_results,
            "Product Name":   product,
            "Price":          price,
        }
        for v in validations:
            row = dict(base)
            row["Step"]     = v.get("step", "")
            row["Expected"] = v.get("expected", "")
            row["Actual"]   = v.get("actual", "")
            row["Status"]   = v.get("status", "")
            row["Error"]    = v.get("error", "")
            self.rows.append(row)
        self._write()

    def export(self):
        self._write()
        print(f"[Report saved] {self.path}")
        return self.path

    def _write(self):
        if not self.rows:
            return

        df = pd.DataFrame(self.rows, columns=COLUMNS)

        with pd.ExcelWriter(self.path, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Results")
            ws = writer.sheets["Results"]
            self._format_sheet(ws, df)

        # Also write a Summary sheet
        self._write_summary(df)

    def _format_sheet(self, ws, df):
        # Header row styling
        for col_idx, col_name in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = BLUE_HEADER
            cell.font = WHITE_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = THIN_BORDER

        status_col   = COLUMNS.index("Status") + 1
        overall_col  = COLUMNS.index("Overall Status") + 1
        noresult_col = COLUMNS.index("No Results") + 1

        prev_combo = None
        row_fill = None

        for row_idx in range(2, len(df) + 2):
            combo = (
                ws.cell(row=row_idx, column=1).value,   # Browser
                ws.cell(row=row_idx, column=2).value,   # Q1
                ws.cell(row=row_idx, column=3).value,   # Q2
                ws.cell(row=row_idx, column=4).value,   # Q3
                ws.cell(row=row_idx, column=5).value,   # Q4
                ws.cell(row=row_idx, column=6).value,   # Q5
                ws.cell(row=row_idx, column=7).value,   # Q6
            )

            # Alternate background per test combo for readability
            if combo != prev_combo:
                prev_combo = combo
                row_fill = GREY if (row_idx % 2 == 0) else None

            status_val  = ws.cell(row=row_idx, column=status_col).value
            overall_val = ws.cell(row=row_idx, column=overall_col).value
            noresult    = ws.cell(row=row_idx, column=noresult_col).value

            for col_idx in range(1, len(COLUMNS) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = NORMAL_FONT
                cell.border = THIN_BORDER
                cell.alignment = Alignment(vertical="center", wrap_text=False)

                # Color by status
                if col_idx == status_col:
                    if status_val == "PASS":
                        cell.fill = GREEN
                    elif status_val == "FAIL":
                        cell.fill = RED
                    elif status_val == "SKIP":
                        cell.fill = YELLOW
                elif col_idx == overall_col:
                    if overall_val == "PASS":
                        cell.fill = GREEN
                        cell.font = Font(bold=True, color="375623", size=10)
                    elif overall_val == "FAIL":
                        cell.fill = RED
                        cell.font = Font(bold=True, color="9C0006", size=10)
                elif col_idx == noresult_col:
                    if noresult == "YES":
                        cell.fill = ORANGE
                        cell.font = Font(bold=True, size=10)
                elif row_fill:
                    cell.fill = row_fill

        # Column widths
        col_widths = {
            "Browser": 12, "Q1": 5, "Q2": 5, "Q3": 5, "Q4": 5, "Q5": 5, "Q6": 5,
            "Overall Status": 14, "No Results": 12, "Product Name": 45, "Price": 10,
            "Step": 45, "Expected": 40, "Actual": 55, "Status": 10, "Error": 40,
        }
        for col_idx, col_name in enumerate(COLUMNS, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = col_widths.get(col_name, 20)

        # Freeze header + first 7 columns
        ws.freeze_panes = "H2"
        ws.row_dimensions[1].height = 30

    def _write_summary(self, df):
        """Add a Summary sheet with one row per test combination."""
        try:
            summary_rows = []
            grouped = df.groupby(["Browser", "Q1", "Q2", "Q3", "Q4", "Q5", "Q6"])
            for keys, group in grouped:
                browser, q1, q2, q3, q4, q5, q6 = keys
                overall   = group["Overall Status"].iloc[0]
                no_result = group["No Results"].iloc[0]
                product   = group["Product Name"].iloc[0]
                price     = group["Price"].iloc[0]
                total     = len(group)
                passed    = (group["Status"] == "PASS").sum()
                failed    = (group["Status"] == "FAIL").sum()
                skipped   = (group["Status"] == "SKIP").sum()
                failed_steps = " | ".join(group[group["Status"] == "FAIL"]["Step"].tolist())
                summary_rows.append({
                    "Browser": browser, "Q1": q1, "Q2": q2, "Q3": q3,
                    "Q4": q4, "Q5": q5, "Q6": q6,
                    "Overall Status": overall, "No Results": no_result,
                    "Product Name": product, "Price": price,
                    "Total Steps": total, "Passed": passed,
                    "Failed": failed, "Skipped": skipped,
                    "Failed Steps": failed_steps,
                })

            summary_df = pd.DataFrame(summary_rows)

            with pd.ExcelWriter(self.path, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
                summary_df.to_excel(writer, index=False, sheet_name="Summary")
                ws = writer.sheets["Summary"]

                sum_cols = list(summary_df.columns)
                overall_col  = sum_cols.index("Overall Status") + 1
                noresult_col = sum_cols.index("No Results") + 1
                failed_col   = sum_cols.index("Failed") + 1

                # Header
                for col_idx in range(1, len(sum_cols) + 1):
                    cell = ws.cell(row=1, column=col_idx)
                    cell.fill = BLUE_HEADER
                    cell.font = WHITE_FONT
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.border = THIN_BORDER

                for row_idx in range(2, len(summary_df) + 2):
                    overall_val  = ws.cell(row=row_idx, column=overall_col).value
                    noresult_val = ws.cell(row=row_idx, column=noresult_col).value
                    failed_val   = ws.cell(row=row_idx, column=failed_col).value

                    for col_idx in range(1, len(sum_cols) + 1):
                        cell = ws.cell(row=row_idx, column=col_idx)
                        cell.font = NORMAL_FONT
                        cell.border = THIN_BORDER
                        cell.alignment = Alignment(vertical="center")

                        if col_idx == overall_col:
                            cell.fill = GREEN if overall_val == "PASS" else RED
                            cell.font = BOLD_FONT
                        elif col_idx == noresult_col and noresult_val == "YES":
                            cell.fill = ORANGE
                            cell.font = BOLD_FONT
                        elif col_idx == failed_col and failed_val and int(failed_val) > 0:
                            cell.fill = RED

                # Column widths for summary
                summary_widths = {
                    "Browser": 12, "Q1": 5, "Q2": 5, "Q3": 5, "Q4": 5, "Q5": 5, "Q6": 5,
                    "Overall Status": 14, "No Results": 12, "Product Name": 45, "Price": 10,
                    "Total Steps": 12, "Passed": 10, "Failed": 10, "Skipped": 10,
                    "Failed Steps": 60,
                }
                for col_idx, col_name in enumerate(sum_cols, start=1):
                    ws.column_dimensions[get_column_letter(col_idx)].width = summary_widths.get(col_name, 15)

                ws.freeze_panes = "A2"
                ws.row_dimensions[1].height = 30

        except Exception as e:
            print(f"[Summary sheet error] {e}")
