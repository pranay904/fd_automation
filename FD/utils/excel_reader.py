from openpyxl import load_workbook


def read_excel(file_path, sheet_name):
    """
    Reads Excel file and returns list of dictionaries
    """
    wb = load_workbook(filename=file_path)
    sheet = wb[sheet_name]

    headers = [cell.value for cell in sheet[1]]
    data = []

    for row in sheet.iter_rows(min_row=2, values_only=True):
        row_dict = dict(zip(headers, row))
        data.append(row_dict)

    wb.close()
    return data
